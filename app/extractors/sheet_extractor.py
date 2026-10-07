"""Extractor for spreadsheets (Excel .xlsx and CSV .csv files)."""

import csv
import io
import re
from typing import Optional
import openpyxl

from app.extractors.base import BaseExtractor, ExtractedQuestion, ExtractionResult
from app.utils.logger import setup_logger

logger = setup_logger(__name__)

QUESTION_HEADER_RE = re.compile(
    r"^(?:question[_\s\-]*text|question|questions|q[_\s\-]*text|prompt|question[_\s\-]*(?:\(english\)|english)|english[_\s\-]*question|question[_\s\-]*title|सवाल|प्रश्न)$",
    re.IGNORECASE,
)
HINDI_QUESTION_HEADER_RE = re.compile(
    r"^(?:hindi[_\s\-]*question|question[_\s\-]*(?:\(hindi\)|hindi)|q[_\s\-]*hindi|hindi[_\s\-]*q|question[_\s\-]*hin|hindi|हिंदी|हिंदी[_\s\-]*प्रश्न|सवाल[_\s\-]*\(हिंदी\)|प्रश्न[_\s\-]*\(हिंदी\))$",
    re.IGNORECASE,
)
OPTION_HINDI_HEADER_RE = re.compile(
    r"^(?:(?:option|opt|choice|विकल्प)[_\s\-]*([a-z0-9]+)[_\s\-]*(?:hindi|hin)|\(([a-z0-9]+)\)[_\s\-]*(?:hindi|hin)|(?:hindi|hin)[_\s\-]*(?:option|opt|choice)[_\s\-]*([a-z0-9]+))$",
    re.IGNORECASE,
)
ANSWER_HEADER_RE = re.compile(
    r"^(?:answer|ans|correct[_\s\-]*answer|correct[_\s\-]*option|correct|key|right[_\s\-]*answer|उत्तर|सही[_\s\-]*उत्तर)$",
    re.IGNORECASE,
)
EXPLANATION_HEADER_RE = re.compile(
    r"^(?:explanation|expl|note|notes|why|reason|व्याख्या|स्पष्टीकरण)$",
    re.IGNORECASE,
)
QNUM_HEADER_RE = re.compile(
    r"^(?:q[_\s\-]*no|qnum|q[_\s\-]*num|qid|q_id|id|no|s\.?no|sr\.?no|क्रम[_\s\-]*संख्या|#)$",
    re.IGNORECASE,
)
OPTION_HEADER_RE = re.compile(
    r"^(?:(?:option|opt|choice|विकल्प)[_\s\-]*([a-z0-9]+)|([a-e])|\(([a-z0-9]+)\)|\[([a-z0-9]+)\])$",
    re.IGNORECASE,
)


class SheetExtractor(BaseExtractor):
    """Extracts questions, options, answers, and explanations from Excel and CSV sheets."""

    def extract_from_bytes(self, content: bytes, filename: str) -> ExtractionResult:
        """Parse raw bytes of an .xlsx or .csv file."""
        lower_name = filename.lower()
        if lower_name.endswith(".csv"):
            return self._extract_csv(content, filename)
        elif lower_name.endswith((".xlsx", ".xlsm", ".xltx")):
            return self._extract_excel(content, filename)
        else:
            return ExtractionResult(
                error_message=f"Unsupported spreadsheet format: {filename}. Please upload .xlsx or .csv files.",
                source_filename=filename,
            )

    def extract_from_file(self, file_path: str) -> ExtractionResult:
        """Read and parse from local file path."""
        with open(file_path, "rb") as f:
            content = f.read()
        return self.extract_from_bytes(content, filename=file_path)

    def _extract_csv(self, content: bytes, filename: str) -> ExtractionResult:
        """Decode and parse CSV content."""
        # Try UTF-8 with BOM, then standard UTF-8, then fallback to cp1252 / latin-1
        text = None
        for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                text = content.decode(enc)
                break
            except UnicodeDecodeError:
                continue

        if text is None:
            return ExtractionResult(
                error_message="Could not decode CSV file. Please ensure it is saved in UTF-8 encoding.",
                source_filename=filename,
            )

        # Detect delimiter (comma, semicolon, tab)
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
            delimiter = dialect.delimiter
        except Exception:
            delimiter = ","

        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        raw_rows = [row for row in reader]
        return self._process_grid(raw_rows, filename)

    def _extract_excel(self, content: bytes, filename: str) -> ExtractionResult:
        """Parse Excel workbook with openpyxl."""
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
        except Exception as e:
            logger.exception("Failed to open Excel workbook %s: %s", filename, e)
            return ExtractionResult(
                error_message=f"Failed to read Excel file: {str(e)}",
                source_filename=filename,
            )

        best_result: Optional[ExtractionResult] = None

        # Check each non-empty sheet in the workbook
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_in_sheet: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                str_row = [str(cell).strip() if cell is not None else "" for cell in row]
                if any(str_row):
                    rows_in_sheet.append(str_row)

            if not rows_in_sheet:
                continue

            result = self._process_grid(rows_in_sheet, filename)
            # If this sheet yielded valid questions, return it immediately
            if result.success and len(result.questions) > 0:
                logger.info("Successfully extracted %d questions from sheet '%s' in %s", len(result.questions), sheet_name, filename)
                return result

            # Track best result across sheets
            if best_result is None or (result.total_found > best_result.total_found):
                best_result = result

        if best_result is not None:
            return best_result

        return ExtractionResult(
            error_message="The Excel file is empty or contains no readable sheets.",
            source_filename=filename,
        )

    def _process_grid(self, raw_rows: list[list[str]], filename: str) -> ExtractionResult:
        """Convert a 2D grid of rows into structured ExtractedQuestions."""
        # Clean rows
        cleaned_rows: list[list[str]] = []
        for row in raw_rows:
            cleaned = [str(c).strip() for c in row]
            while cleaned and cleaned[-1] == "":
                cleaned.pop()
            if any(cleaned):
                cleaned_rows.append(cleaned)

        if not cleaned_rows:
            return ExtractionResult(
                error_message="No readable data rows found in the document.",
                source_filename=filename,
            )

        # Check if single-column raw text sheet
        max_cols = max(len(r) for r in cleaned_rows)
        if max_cols == 1:
            raw_combined = "\n".join(r[0] for r in cleaned_rows if r)
            return ExtractionResult(
                raw_text=raw_combined,
                total_found=1,
                source_filename=filename,
            )

        # Try to detect header row
        header_row_idx, col_mapping = self._detect_headers(cleaned_rows)

        questions: list[ExtractedQuestion] = []
        if header_row_idx is not None and col_mapping.get("question") is not None:
            # Header-driven extraction
            data_rows = cleaned_rows[header_row_idx + 1 :]
            for idx, row in enumerate(data_rows, start=1):
                extracted = self._extract_from_mapped_row(row, col_mapping, default_number=idx)
                if extracted:
                    questions.append(extracted)
        else:
            # Heuristic-based extraction without recognized headers
            for idx, row in enumerate(cleaned_rows, start=1):
                extracted = self._extract_from_heuristic_row(row, default_number=idx)
                if extracted:
                    questions.append(extracted)

        if not questions:
            return ExtractionResult(
                error_message="Could not detect questions in the sheet. Please ensure columns include Question, Options (A, B, C, D), and Answer.",
                source_filename=filename,
            )

        # Generate formatted raw text compatible with QuizBotProParser
        formatted_blocks = [q.to_formatted_block(default_number=i) for i, q in enumerate(questions, start=1)]
        raw_text = "\n\n".join(formatted_blocks)

        return ExtractionResult(
            questions=questions,
            raw_text=raw_text,
            total_found=len(questions),
            source_filename=filename,
        )

    def _detect_headers(self, rows: list[list[str]]) -> tuple[Optional[int], dict]:
        """Inspect rows to identify header mapping (scanning up to first 50 rows to bypass dashboards)."""
        max_scan = min(50, len(rows))
        for r_idx in range(max_scan):
            row = rows[r_idx]
            mapping = {
                "question": None,
                "hindi_question": None,
                "options": [],  # list of (col_idx, letter_or_key)
                "options_hindi": {},  # key -> col_idx
                "answer": None,
                "explanation": None,
                "q_num": None,
            }

            for c_idx, cell in enumerate(row):
                cell_clean = cell.strip()
                if not cell_clean:
                    continue

                if HINDI_QUESTION_HEADER_RE.match(cell_clean):
                    mapping["hindi_question"] = c_idx
                elif QUESTION_HEADER_RE.match(cell_clean):
                    mapping["question"] = c_idx
                elif ANSWER_HEADER_RE.match(cell_clean):
                    mapping["answer"] = c_idx
                elif EXPLANATION_HEADER_RE.match(cell_clean):
                    mapping["explanation"] = c_idx
                elif QNUM_HEADER_RE.match(cell_clean):
                    mapping["q_num"] = c_idx
                else:
                    opt_hin_m = OPTION_HINDI_HEADER_RE.match(cell_clean)
                    if opt_hin_m:
                        key = next((g for g in opt_hin_m.groups() if g is not None), "").upper()
                        mapping["options_hindi"][key] = c_idx
                    else:
                        opt_m = OPTION_HEADER_RE.match(cell_clean)
                        if opt_m:
                            key = next((g for g in opt_m.groups() if g is not None), "").upper()
                            mapping["options"].append((c_idx, key))

            # If only hindi question column was found, treat as question
            if mapping["question"] is None and mapping["hindi_question"] is not None:
                mapping["question"] = mapping["hindi_question"]

            # If we found at least a question column and (at least 2 option columns OR an answer column)
            if mapping["question"] is not None and (len(mapping["options"]) >= 2 or mapping["answer"] is not None):
                # Sort options by column index
                mapping["options"].sort(key=lambda x: x[0])
                return r_idx, mapping

        return None, {}

    def _extract_from_mapped_row(
        self, row: list[str], mapping: dict, default_number: int
    ) -> Optional[ExtractedQuestion]:
        """Extract a single question using the detected header column map."""
        q_idx = mapping.get("question")
        if q_idx is None or q_idx >= len(row):
            return None

        question_text = row[q_idx].strip()
        if not question_text:
            return None

        # Ignore summary / metadata rows (e.g. "Total", "Average", "Note:")
        if re.match(r"^(?:total|average|sum|grand\s*total|note|remarks?)[:\s]", question_text, re.IGNORECASE):
            return None

        # Append separate Hindi question column if present
        hindi_idx = mapping.get("hindi_question")
        if hindi_idx is not None and hindi_idx != q_idx and hindi_idx < len(row):
            hindi_text = row[hindi_idx].strip()
            if hindi_text:
                if not question_text.endswith("?"):
                    question_text = f"{question_text}? {hindi_text}"
                else:
                    question_text = f"{question_text} {hindi_text}"

        options: list[str] = []
        options_hindi = mapping.get("options_hindi", {})
        for c_idx, key in mapping.get("options", []):
            if c_idx < len(row):
                val = row[c_idx].strip()
                if val:
                    # Check if there's a Hindi translation column for this option
                    if key in options_hindi and options_hindi[key] < len(row):
                        hin_val = row[options_hindi[key]].strip()
                        if hin_val:
                            val = f"{val} ({hin_val})"
                    options.append(val)

        # If options weren't in dedicated columns, check for standard remaining columns
        if not options and len(row) > q_idx + 1:
            for extra_idx in range(q_idx + 1, len(row)):
                if (
                    extra_idx != mapping.get("answer")
                    and extra_idx != mapping.get("explanation")
                    and extra_idx != mapping.get("hindi_question")
                ):
                    val = row[extra_idx].strip()
                    if val:
                        options.append(val)

        ans_idx = mapping.get("answer")
        answer_raw = row[ans_idx].strip() if ans_idx is not None and ans_idx < len(row) else None
        if answer_raw == "":
            answer_raw = None

        expl_idx = mapping.get("explanation")
        explanation = row[expl_idx].strip() if expl_idx is not None and expl_idx < len(row) else None
        if explanation == "":
            explanation = None

        qnum_idx = mapping.get("q_num")
        q_num = None
        if qnum_idx is not None and qnum_idx < len(row):
            num_str = row[qnum_idx].strip()
            if num_str.isdigit():
                q_num = int(num_str)

        return ExtractedQuestion(
            question_text=question_text,
            options=options,
            answer_raw=answer_raw,
            explanation=explanation,
            question_number=q_num or default_number,
        )

    def _extract_from_heuristic_row(
        self, row: list[str], default_number: int
    ) -> Optional[ExtractedQuestion]:
        """Heuristically parse a row without headers (e.g. Col 0: Question, Col 1..4: Options, Col 5: Answer)."""
        clean_cells = [c.strip() for c in row if c.strip()]
        if len(clean_cells) < 3:
            return None

        question_text = clean_cells[0]
        # Check if question text is just a header row like "Question"
        if QUESTION_HEADER_RE.match(question_text):
            return None

        remaining = clean_cells[1:]
        options: list[str] = []
        answer_raw: Optional[str] = None
        explanation: Optional[str] = None

        # Check if last cell looks like an answer (e.g. 'A', 'B', '1', '2', or marked with ✅)
        last_cell = remaining[-1]
        is_answer = False
        if len(last_cell) <= 15 and (
            last_cell.upper() in ("A", "B", "C", "D", "E", "1", "2", "3", "4", "5")
            or any(m in last_cell for m in ("✅", "✔️", "✔", "*"))
            or re.match(r"^(?:Ans|Answer|उत्तर)[\:\-\s]", last_cell, re.IGNORECASE)
        ):
            is_answer = True

        if is_answer and len(remaining) >= 3:
            answer_raw = last_cell
            options = remaining[:-1]
        else:
            options = remaining

        return ExtractedQuestion(
            question_text=question_text,
            options=options,
            answer_raw=answer_raw,
            explanation=explanation,
            question_number=default_number,
        )


def generate_quiz_template_bytes() -> bytes:
    """Generate a styled Excel workbook (.xlsx) template for quiz question creation."""
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quiz Questions"

    headers = ["Question", "Option A", "Option B", "Option C", "Option D", "Answer", "Explanation"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Sample rows demonstrating different subjects and bilingual questions
    ws.append([
        "What is the capital of India? भारत की राजधानी क्या है?",
        "Mumbai", "New Delhi", "Kolkata", "Chennai",
        "B", "New Delhi is the official capital of India."
    ])
    ws.append([
        "Which planet is known as the Red Planet?",
        "Venus", "Mars", "Jupiter", "Saturn",
        "B", "Mars appears red due to iron oxide (rust) on its surface."
    ])
    ws.append([
        "How many types of North are there? उत्तर कितने प्रकार के होते हैं?",
        "3", "5", "8", "10",
        "A", "True North, Grid North, and Magnetic North."
    ])

    # Adjust column widths
    column_widths = [45, 18, 18, 18, 18, 12, 35]
    for i, width in enumerate(column_widths, start=1):
        col_letter = openpyxl.utils.get_column_letter(i)
        ws.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

