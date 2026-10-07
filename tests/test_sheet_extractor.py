"""Unit tests for Phase 1 spreadsheet (Excel/CSV) question extractor."""

import io
import pytest
import openpyxl

from app.extractors.sheet_extractor import SheetExtractor, generate_quiz_template_bytes
from app.parser.parser import BulkQuizParser


def test_csv_extraction_with_standard_headers():
    csv_content = """Question,Option A,Option B,Option C,Option D,Answer,Explanation
What is the capital of France?,Paris,London,Berlin,Madrid,A,Paris is the capital of France.
Which planet is known as the Red Planet?,Venus,Mars,Jupiter,Saturn,B,Mars appears red due to iron oxide.
"""
    file_bytes = csv_content.encode("utf-8")
    extractor = SheetExtractor()
    result = extractor.extract(file_bytes, "questions.csv")

    assert result.success is True
    assert result.error is None
    assert len(result.questions) == 2

    q1 = result.questions[0]
    assert "Paris" in q1.options
    assert q1.correct_index == 0
    assert q1.explanation == "Paris is the capital of France."

    # Verify formatting can be parsed cleanly by BulkQuizParser
    block = result.to_formatted_text()
    parse_res = BulkQuizParser().parse(block)
    assert not parse_res.has_errors
    assert parse_res.valid_count == 2
    assert parse_res.questions[0].correct_option == 0
    assert parse_res.questions[1].correct_option == 1


def test_csv_extraction_bilingual_hindi_english():
    csv_content = """Question,Hindi Question,Option A,Option B,Option C,Option D,Answer
What is the motto of NCC?,NCC का आदर्श वाक्य क्या है?,Unity and Duty,Unity and Discipline,Discipline and Service,Unity and Bravery,B
"""
    file_bytes = csv_content.encode("utf-8-sig")  # Test with UTF-8 BOM
    extractor = SheetExtractor()
    result = extractor.extract(file_bytes, "ncc_quiz.csv")

    assert result.success is True
    assert len(result.questions) == 1
    q = result.questions[0]
    assert "What is the motto of NCC?" in q.question_text
    assert "NCC का आदर्श वाक्य क्या है?" in q.question_text
    assert q.correct_index == 1

    formatted_block = q.to_formatted_block()
    assert "✅" in formatted_block
    assert "Unity and Discipline" in formatted_block


def test_excel_xlsx_extraction():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quiz Questions"

    ws.append(["Question", "A", "B", "C", "D", "Correct Answer", "Explanation"])
    ws.append(["What is 2 + 2?", "3", "4", "5", "6", "B", "2+2=4"])
    ws.append(["What is the speed of light?", "3x10^8 m/s", "150 m/s", "1000 km/h", "Sound speed", "A", ""])

    out = io.BytesIO()
    wb.save(out)
    excel_bytes = out.getvalue()

    extractor = SheetExtractor()
    result = extractor.extract(excel_bytes, "math.xlsx")

    assert result.success is True
    assert len(result.questions) == 2
    assert result.questions[0].correct_index == 1
    assert result.questions[0].explanation == "2+2=4"
    assert result.questions[1].correct_index == 0


def test_generate_quiz_template_bytes():
    template_data = generate_quiz_template_bytes()
    assert len(template_data) > 0

    # Ensure it is a valid Excel workbook that openpyxl can load
    wb = openpyxl.load_workbook(io.BytesIO(template_data), data_only=True)
    assert "Quiz Questions" in wb.sheetnames
    ws = wb["Quiz Questions"]
    assert ws.cell(row=1, column=1).value == "Question"

    # Test extracting directly from the generated template!
    extractor = SheetExtractor()
    result = extractor.extract(template_data, "sample_template.xlsx")
    assert result.success is True
    assert len(result.questions) >= 3  # The 3 sample rows in the template


def test_single_column_raw_text_excel():
    wb = openpyxl.Workbook()
    ws = wb.active

    # Single column with complete question text blocks
    q_block = """Q1. What is Python?
A. A snake
B. A programming language ✅
C. A car
D. An island"""
    ws.append([q_block])

    out = io.BytesIO()
    wb.save(out)
    excel_bytes = out.getvalue()

    extractor = SheetExtractor()
    result = extractor.extract(excel_bytes, "raw_text.xlsx")

    assert result.success is True
    assert "What is Python?" in result.to_formatted_text()
    parse_res = BulkQuizParser().parse(result.to_formatted_text())
    assert not parse_res.has_errors
    assert parse_res.valid_count == 1
    assert parse_res.questions[0].correct_option == 1


def test_unsupported_file_extension():
    extractor = SheetExtractor()
    result = extractor.extract(b"dummy data", "document.pdf")
    assert result.success is False
    assert "Unsupported spreadsheet format" in result.error


def test_spreadsheet_with_dashboard_and_underscored_headers():
    """Test extracting from a spreadsheet like 'Untitled spreadsheet.xlsx' with summary dashboard and Option_A headers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NCC Quiz"

    # Row 1-7: Summary Dashboard
    ws.append(["Summary Dashboard"])
    ws.append(["Quiz Name", "NCC Practice Quiz 16", "Topic Distribution", "Map Reading", 6])
    ws.append(["Total Questions", 25, "Difficulty Distribution", "Easy", 10])
    ws.append(["New Questions", 17, "", "Medium", 15])
    ws.append(["Revision Questions", 8, "", "Hard", 0])
    ws.append(["Quiz Creation Date", "2026-06-01"])
    ws.append(["Average Times_Used", 0.32])
    ws.append([])  # Blank row 8
    ws.append([])  # Blank row 9

    # Row 10: Column headers with underscores
    ws.append([
        "Quiz_No", "QID", "Question", "Option_A", "Option_B", "Option_C", "Option_D",
        "Correct_Answer", "Topic", "Difficulty", "New_or_Revision", "Marks"
    ])

    # Row 11: Question with numeric options and numeric answer text ("360")
    ws.append([
        "Quiz 16", "NCC601", "How many degrees are there in a compass?",
        "360", "270", "180", "90", "360", "Map Reading", "Easy", "New", 1
    ])

    # Row 12: Question with string options and string answer text ("North")
    ws.append([
        "Quiz 16", "NCC602", "Which direction is indicated by the top of a map?",
        "North", "West", "South", "East", "North", "Map Reading", "Easy", "New", 1
    ])

    # Row 13: Question with answer matching Option C ("Platoon Commander")
    ws.append([
        "Quiz 16", "NCC605", "Whose orders will the Section Commander wait for after reorganization in Section Battle Drill?",
        "Commanding Officer", "Company Commander", "Platoon Commander", "None of the above",
        "Platoon Commander", "Field Craft & Battle Craft", "Easy", "New", 1
    ])

    out = io.BytesIO()
    wb.save(out)
    excel_bytes = out.getvalue()

    extractor = SheetExtractor()
    result = extractor.extract(excel_bytes, "Untitled spreadsheet.xlsx")

    assert result.success is True
    assert len(result.questions) == 3

    # Check question 1: 360 should be resolved to Option A (index 0)
    q1 = result.questions[0]
    assert q1.question_text == "How many degrees are there in a compass?"
    assert q1.correct_index == 0

    # Check question 2: North should be resolved to Option A (index 0)
    q2 = result.questions[1]
    assert q2.question_text == "Which direction is indicated by the top of a map?"
    assert q2.correct_index == 0

    # Check question 3: Platoon Commander should be resolved to Option C (index 2)
    q3 = result.questions[2]
    assert q3.correct_index == 2

    # Verify BulkQuizParser parses the formatted blocks into valid questions
    parse_res = BulkQuizParser().parse(result.to_formatted_text())
    assert not parse_res.has_errors
    assert parse_res.valid_count == 3
    assert parse_res.questions[0].correct_option == 0
    assert parse_res.questions[1].correct_option == 0
    assert parse_res.questions[2].correct_option == 2


def test_bilingual_options_in_separate_columns():
    """Test extracting when options and Hindi translations are in separate columns."""
    csv_content = """Question,Option_A,Option_A_Hindi,Option_B,Option_B_Hindi,Correct_Answer
What is the motto of NCC?,Unity and Duty,एकता और कर्तव्य,Unity and Discipline,एकता और अनुशासन,B
"""
    file_bytes = csv_content.encode("utf-8")
    extractor = SheetExtractor()
    result = extractor.extract(file_bytes, "bilingual_options.csv")

    assert result.success is True
    assert len(result.questions) == 1
    q = result.questions[0]
    assert "Unity and Duty (एकता और कर्तव्य)" in q.options[0]
    assert "Unity and Discipline (एकता और अनुशासन)" in q.options[1]
    assert q.correct_index == 1


def test_multi_sheet_excel_workbook_picks_questions_sheet():
    """Test that a workbook with a non-question cover sheet correctly finds the questions sheet."""
    wb = openpyxl.Workbook()
    # Sheet 1: Cover Page
    ws1 = wb.active
    ws1.title = "Cover Page"
    ws1.append(["Quiz Application Details"])
    ws1.append(["Author", "Instructor"])

    # Sheet 2: Questions
    ws2 = wb.create_sheet(title="Quiz Questions")
    ws2.append(["Question", "Option A", "Option B", "Answer"])
    ws2.append(["Is the Earth round?", "Yes", "No", "A"])

    out = io.BytesIO()
    wb.save(out)
    excel_bytes = out.getvalue()

    extractor = SheetExtractor()
    result = extractor.extract(excel_bytes, "multi_sheet.xlsx")

    assert result.success is True
    assert len(result.questions) == 1
    assert result.questions[0].question_text == "Is the Earth round?"

