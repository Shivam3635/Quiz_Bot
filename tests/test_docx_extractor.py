"""Unit tests for Phase 2 Word document (.docx) question extractor."""

import io
import pytest
import docx

from app.extractors.docx_extractor import DocxExtractor
from app.parser.parser import BulkQuizParser


def test_docx_paragraphs_extraction():
    """Test extracting questions typed as regular text paragraphs in a Word document."""
    doc = docx.Document()
    doc.add_heading("Sample Quiz", level=1)
    doc.add_paragraph("Q1. What is the capital of India? भारत की राजधानी क्या है?")
    doc.add_paragraph("A. Mumbai")
    doc.add_paragraph("B. New Delhi ✅")
    doc.add_paragraph("C. Kolkata")
    doc.add_paragraph("D. Chennai")
    doc.add_paragraph("Explanation: New Delhi is the capital.")

    doc.add_paragraph("Q2. Which planet is closest to the Sun?")
    doc.add_paragraph("A) Mercury")
    doc.add_paragraph("B) Venus")
    doc.add_paragraph("C) Mars")
    doc.add_paragraph("D) Jupiter")
    doc.add_paragraph("Answer: A")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    extractor = DocxExtractor()
    result = extractor.extract(docx_bytes, "science_quiz.docx")

    assert result.success is True
    assert result.error is None
    assert len(result.questions) == 2

    # Check question 1
    q1 = result.questions[0]
    assert "What is the capital of India?" in q1.question_text
    assert q1.correct_index == 1
    assert q1.explanation == "New Delhi is the capital."

    # Check question 2
    q2 = result.questions[1]
    assert "Mercury" in q2.options
    assert q2.correct_index == 0

    # Parse with BulkQuizParser
    parse_res = BulkQuizParser().parse(result.to_formatted_text())
    assert not parse_res.has_errors
    assert parse_res.valid_count == 2
    assert parse_res.questions[0].correct_option == 1
    assert parse_res.questions[1].correct_option == 0


def test_docx_table_extraction():
    """Test extracting questions from a Word document containing a structured table."""
    doc = docx.Document()
    doc.add_heading("NCC Questions Table", level=1)

    table = doc.add_table(rows=1, cols=6)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Question"
    hdr_cells[1].text = "Option A"
    hdr_cells[2].text = "Option B"
    hdr_cells[3].text = "Option C"
    hdr_cells[4].text = "Option D"
    hdr_cells[5].text = "Correct Answer"

    # Add question 1
    row1 = table.add_row().cells
    row1[0].text = "How many degrees are there in a compass?"
    row1[1].text = "360"
    row1[2].text = "270"
    row1[3].text = "180"
    row1[4].text = "90"
    row1[5].text = "360"

    # Add question 2
    row2 = table.add_row().cells
    row2[0].text = "Which direction is indicated by the top of a map?"
    row2[1].text = "North"
    row2[2].text = "West"
    row2[3].text = "South"
    row2[4].text = "East"
    row2[5].text = "North"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    extractor = DocxExtractor()
    result = extractor.extract(docx_bytes, "ncc_table.docx")

    assert result.success is True
    assert len(result.questions) == 2
    assert result.questions[0].correct_index == 0
    assert result.questions[1].correct_index == 0

    parse_res = BulkQuizParser().parse(result.to_formatted_text())
    assert not parse_res.has_errors
    assert parse_res.valid_count == 2


def test_docx_combined_paragraphs_and_table():
    """Test Word document with both paragraphs and a table."""
    doc = docx.Document()

    # Paragraph question
    doc.add_paragraph("Q1. What is Python?")
    doc.add_paragraph("A. A programming language ✅")
    doc.add_paragraph("B. A car")

    # Table question
    table = doc.add_table(rows=1, cols=4)
    table.rows[0].cells[0].text = "Question"
    table.rows[0].cells[1].text = "Option A"
    table.rows[0].cells[2].text = "Option B"
    table.rows[0].cells[3].text = "Answer"

    r = table.add_row().cells
    r[0].text = "Is 2 an even number?"
    r[1].text = "Yes"
    r[2].text = "No"
    r[3].text = "A"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    extractor = DocxExtractor()
    result = extractor.extract(docx_bytes, "mixed.docx")

    assert result.success is True
    assert len(result.questions) == 2

    parse_res = BulkQuizParser().parse(result.to_formatted_text())
    assert not parse_res.has_errors
    assert parse_res.valid_count == 2


def test_docx_unsupported_extension():
    extractor = DocxExtractor()
    result = extractor.extract(b"dummy data", "document.pdf")
    assert result.success is False
    assert "Unsupported Word document format" in result.error


def test_docx_bulleted_options_with_unicode_checkmark():
    """Test Word document with bulleted options (• A.) and ✓ checkmark."""
    doc = docx.Document()
    doc.add_paragraph("NCC Practice Quiz 16")
    doc.add_paragraph("25 MCQs • English + Hindi • Correct answers marked with ✓")
    doc.add_paragraph("Q1. What is the effective range of a .22 Deluxe Rifle? .22 डीलक्स राइफल की कारगर रेंज क्या है?")
    doc.add_paragraph("• A. 25 Yards (25 गज) ✓")
    doc.add_paragraph("• B. 50 Yards (50 गज)")
    doc.add_paragraph("• C. 300 Yards (300 गज)")
    doc.add_paragraph("• D. 1700 Yards (1700 गज)")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    extractor = DocxExtractor()
    result = extractor.extract(docx_bytes, "NCC_Practice_Quiz_16.docx")

    assert result.success is True
    assert len(result.questions) == 1
    assert result.questions[0].correct_index == 0
    assert len(result.questions[0].options) == 4

    parse_res = BulkQuizParser().parse(result.to_formatted_text())
    assert not parse_res.has_errors
    assert parse_res.valid_count == 1
    assert parse_res.questions[0].correct_option == 0
